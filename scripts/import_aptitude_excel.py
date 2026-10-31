"""
StudyFlow AI - Excel Aptitude Question Bank Importer
Imports, validates, normalizes, and idempotently populates SQLite from:
1. placement_aptitude_800_questions.xlsx (Primary)
2. Placement_Aptitude_Question_Bank_300.xlsx (Secondary)
"""

import os
import re
import json
import sqlite3
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, 'data')
DB_PATH = os.path.join(DATA_DIR, 'studyflow.db')

PRIMARY_FILE = os.path.join(DATA_DIR, 'placement_aptitude_800_questions.xlsx')
SECONDARY_FILE = os.path.join(DATA_DIR, 'Placement_Aptitude_Question_Bank_300.xlsx')
REPORT_FILE = os.path.join(DATA_DIR, 'aptitude_import_report.json')
MASTER_JSON = os.path.join(DATA_DIR, 'aptitude_questions_master.json')

def clean_str(val):
    if val is None:
        return ''
    s = str(val).strip()
    # Normalize unicode quotes and artifacts
    s = s.replace('\ufffd', '-').replace('\u2013', '-').replace('\u2014', '-').replace('\u2018', "'").replace('\u2019', "'").replace('\u201c', '"').replace('\u201d', '"')
    return s

def normalize_category(cat):
    c = clean_str(cat).lower()
    if 'quant' in c:
        return 'Quantitative Aptitude'
    if 'logic' in c or 'reason' in c:
        return 'Logical Reasoning'
    if 'verbal' in c or 'english' in c:
        return 'Verbal Ability'
    if 'data' in c or 'interpretation' in c or 'di' == c:
        return 'Data Interpretation'
    if 'tech' in c or 'computer' in c or 'program' in c or 'software' in c:
        return 'Technical Aptitude'
    return clean_str(cat) or 'Quantitative Aptitude'

def run_import():
    report = {
        "primary_file": PRIMARY_FILE,
        "secondary_file": SECONDARY_FILE,
        "total_questions_found": 0,
        "successfully_imported": 0,
        "auto_fixed_rows": [],
        "invalid_questions": [],
        "duplicate_questions": [],
        "manual_review_required": [],
        "category_counts": {}
    }

    all_questions = []
    seen_texts = {}

    # 1. Process Primary 800 Questions Workbook
    if os.path.exists(PRIMARY_FILE):
        wb800 = openpyxl.load_workbook(PRIMARY_FILE, data_only=True)
        sheet = wb800['800 Aptitude Questions'] if '800 Aptitude Questions' in wb800.sheetnames else wb800.active
        rows = list(sheet.iter_rows(values_only=True))
        header = rows[0]
        data_rows = rows[1:]
        report["total_questions_found"] += len(data_rows)

        for idx, r in enumerate(data_rows, 1):
            orig_no = r[0] if r[0] is not None else idx
            cat = normalize_category(r[1])
            topic = clean_str(r[2]) or 'General'
            q_text = clean_str(r[3])
            opt_a = clean_str(r[4])
            opt_b = clean_str(r[5])
            opt_c = clean_str(r[6])
            opt_d = clean_str(r[7])
            corr_opt = clean_str(r[8]).upper()
            corr_ans = clean_str(r[9])
            explanation = clean_str(r[10])

            # Validation & auto-fix: check if option is empty due to Python "==" symbol parsed as formula
            is_equality_q = 'checks equality in python' in q_text.lower() or 'double-equals operator' in explanation.lower()
            options = {'A': opt_a, 'B': opt_b, 'C': opt_c, 'D': opt_d}
            
            for letter, val in options.items():
                if val == '' and is_equality_q:
                    options[letter] = '=='
                    if corr_opt == letter or corr_ans == '' or corr_ans.lower() == 'none':
                        corr_opt = letter
                        corr_ans = '=='
                    report["auto_fixed_rows"].append({
                        "source": "800_questions",
                        "original_no": orig_no,
                        "question": q_text,
                        "fix": f"Auto-restored Option {letter} and answer to '=='"
                    })

            opt_a, opt_b, opt_c, opt_d = options['A'], options['B'], options['C'], options['D']

            # Validate question text
            if not q_text:
                report["invalid_questions"].append({"original_no": orig_no, "error": "Missing question text"})
                continue

            # Validate options
            if not all([opt_a, opt_b, opt_c, opt_d]):
                report["manual_review_required"].append({
                    "original_no": orig_no,
                    "error": "Missing option values",
                    "options": options
                })
                continue

            # Validate correct option
            if corr_opt not in ['A', 'B', 'C', 'D']:
                report["invalid_questions"].append({"original_no": orig_no, "error": f"Invalid correct option: '{corr_opt}'"})
                continue

            # Check duplicate question text
            norm_text = re.sub(r'\s+', ' ', q_text.lower().strip())
            is_dup = norm_text in seen_texts
            if is_dup:
                report["duplicate_questions"].append({
                    "original_no": orig_no,
                    "question": q_text,
                    "first_seen_id": seen_texts[norm_text]
                })

            qid = f"apt_800_{orig_no}"
            seen_texts[norm_text] = qid

            q_obj = {
                "id": qid,
                "original_no": int(orig_no) if str(orig_no).isdigit() else idx,
                "category": cat,
                "topic": topic,
                "question_text": q_text,
                "option_a": opt_a,
                "option_b": opt_b,
                "option_c": opt_c,
                "option_d": opt_d,
                "correct_option": corr_opt,
                "correct_answer_text": corr_ans or options.get(corr_opt, ''),
                "explanation": explanation or "Refer to standard formula and rules.",
                "difficulty": "Medium",
                "source": "placement_aptitude_800"
            }
            all_questions.append(q_obj)
            report["category_counts"][cat] = report["category_counts"].get(cat, 0) + 1

    # 2. Process Secondary 300 Questions Workbook
    if os.path.exists(SECONDARY_FILE):
        wb300 = openpyxl.load_workbook(SECONDARY_FILE, data_only=True)
        sheet = wb300['Question Bank'] if 'Question Bank' in wb300.sheetnames else wb300.active
        rows = list(sheet.iter_rows(values_only=True))
        data_rows = rows[1:]
        report["total_questions_found"] += len(data_rows)

        for idx, r in enumerate(data_rows, 1):
            orig_id = clean_str(r[0]) or f"APT_{idx:03d}"
            cat = normalize_category(r[1])
            q_text = clean_str(r[2])
            opt_a = clean_str(r[3])
            opt_b = clean_str(r[4])
            opt_c = clean_str(r[5])
            opt_d = clean_str(r[6])
            corr_opt = clean_str(r[7]).upper()
            corr_ans = clean_str(r[8])
            explanation = clean_str(r[9])

            if not q_text:
                continue

            # Check if this exact question text is already imported from primary file
            norm_text = re.sub(r'\s+', ' ', q_text.lower().strip())
            if norm_text in seen_texts:
                report["duplicate_questions"].append({
                    "original_no": orig_id,
                    "question": q_text,
                    "first_seen_id": seen_texts[norm_text],
                    "note": "Skipping duplicate question from 300 bank"
                })
                continue

            if not all([opt_a, opt_b, opt_c, opt_d]) or corr_opt not in ['A', 'B', 'C', 'D']:
                continue

            qid = f"apt_300_{orig_id.replace(' ', '_')}"
            seen_texts[norm_text] = qid

            q_obj = {
                "id": qid,
                "original_no": idx,
                "category": cat,
                "topic": "General Aptitude",
                "question_text": q_text,
                "option_a": opt_a,
                "option_b": opt_b,
                "option_c": opt_c,
                "option_d": opt_d,
                "correct_option": corr_opt,
                "correct_answer_text": corr_ans or {'A': opt_a, 'B': opt_b, 'C': opt_c, 'D': opt_d}.get(corr_opt, ''),
                "explanation": explanation or "Standard formula application.",
                "difficulty": "Medium",
                "source": "placement_aptitude_300"
            }
            all_questions.append(q_obj)
            report["category_counts"][cat] = report["category_counts"].get(cat, 0) + 1

    report["successfully_imported"] = len(all_questions)

    # 3. Save master JSON
    with open(MASTER_JSON, 'w', encoding='utf-8') as f:
        json.dump(all_questions, f, indent=2, ensure_ascii=False)

    # 4. Save validation report
    with open(REPORT_FILE, 'w', encoding='utf-8') as f:
        json.dump(report, f, indent=2, ensure_ascii=False)

    # 5. Populate SQLite studyflow.db idempotently
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    # Ensure schema migrations
    cur.execute("""
        CREATE TABLE IF NOT EXISTS aptitude_questions (
            id TEXT PRIMARY KEY,
            category TEXT NOT NULL,
            question_text TEXT NOT NULL,
            option_a TEXT NOT NULL,
            option_b TEXT NOT NULL,
            option_c TEXT NOT NULL,
            option_d TEXT NOT NULL,
            correct_option TEXT NOT NULL,
            explanation TEXT,
            difficulty TEXT DEFAULT 'Medium',
            topic TEXT,
            correct_answer_text TEXT,
            original_no INTEGER,
            source TEXT DEFAULT 'placement_aptitude_800'
        )
    """)

    cur.execute("PRAGMA table_info(aptitude_questions)")
    existing_cols = [c[1] for c in cur.fetchall()]
    if 'topic' not in existing_cols:
        cur.execute("ALTER TABLE aptitude_questions ADD COLUMN topic TEXT")
    if 'correct_answer_text' not in existing_cols:
        cur.execute("ALTER TABLE aptitude_questions ADD COLUMN correct_answer_text TEXT")
    if 'original_no' not in existing_cols:
        cur.execute("ALTER TABLE aptitude_questions ADD COLUMN original_no INTEGER")
    if 'source' not in existing_cols:
        cur.execute("ALTER TABLE aptitude_questions ADD COLUMN source TEXT DEFAULT 'placement_aptitude_800'")

    # Ensure assessments table
    cur.execute("""
        CREATE TABLE IF NOT EXISTS aptitude_assessments (
            id TEXT PRIMARY KEY,
            student_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            status TEXT DEFAULT 'in_progress',
            started_at TEXT DEFAULT (datetime('now')),
            expires_at TEXT NOT NULL,
            duration_seconds INTEGER DEFAULT 900,
            questions_json TEXT NOT NULL,
            answers_json TEXT DEFAULT '{}',
            score REAL,
            accuracy_percentage REAL,
            created_at TEXT DEFAULT (datetime('now')),
            updated_at TEXT DEFAULT (datetime('now'))
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_aptitude_assessments_student ON aptitude_assessments(student_id)")

    # Ensure aptitude_attempts extra columns
    cur.execute("PRAGMA table_info(aptitude_attempts)")
    attempt_cols = [c[1] for c in cur.fetchall()]
    if 'unanswered_count' not in attempt_cols:
        cur.execute("ALTER TABLE aptitude_attempts ADD COLUMN unanswered_count INTEGER DEFAULT 0")
    if 'category_scores_json' not in attempt_cols:
        cur.execute("ALTER TABLE aptitude_attempts ADD COLUMN category_scores_json TEXT")
    if 'topic_accuracy_json' not in attempt_cols:
        cur.execute("ALTER TABLE aptitude_attempts ADD COLUMN topic_accuracy_json TEXT")
    if 'status' not in attempt_cols:
        cur.execute("ALTER TABLE aptitude_attempts ADD COLUMN status TEXT DEFAULT 'Completed'")
    if 'assessment_id' not in attempt_cols:
        cur.execute("ALTER TABLE aptitude_attempts ADD COLUMN assessment_id TEXT")
    if 'attempt_number' not in attempt_cols:
        cur.execute("ALTER TABLE aptitude_attempts ADD COLUMN attempt_number INTEGER DEFAULT 1")

    # Upsert questions
    upsert_sql = """
        INSERT INTO aptitude_questions (
            id, category, topic, question_text, option_a, option_b, option_c, option_d,
            correct_option, correct_answer_text, explanation, difficulty, original_no, source
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(id) DO UPDATE SET
            category = excluded.category,
            topic = excluded.topic,
            question_text = excluded.question_text,
            option_a = excluded.option_a,
            option_b = excluded.option_b,
            option_c = excluded.option_c,
            option_d = excluded.option_d,
            correct_option = excluded.correct_option,
            correct_answer_text = excluded.correct_answer_text,
            explanation = excluded.explanation,
            difficulty = excluded.difficulty,
            original_no = excluded.original_no,
            source = excluded.source
    """

    for q in all_questions:
        cur.execute(upsert_sql, (
            q["id"], q["category"], q["topic"], q["question_text"],
            q["option_a"], q["option_b"], q["option_c"], q["option_d"],
            q["correct_option"], q["correct_answer_text"], q["explanation"],
            q["difficulty"], q["original_no"], q["source"]
        ))

    conn.commit()
    cur.execute("SELECT COUNT(*) FROM aptitude_questions")
    final_count = cur.fetchone()[0]
    conn.close()

    print("==================================================")
    print("STUDYFLOW AI: APTITUDE QUESTION BANK IMPORT REPORT")
    print("==================================================")
    print(f"Total questions processed:      {report['total_questions_found']}")
    print(f"Successfully imported:         {report['successfully_imported']}")
    print(f"Auto-fixed missing options:    {len(report['auto_fixed_rows'])}")
    print(f"Duplicates detected & handled: {len(report['duplicate_questions'])}")
    print(f"Invalid questions:             {len(report['invalid_questions'])}")
    print(f"Final SQLite question count:   {final_count}")
    print("Category Breakdown:")
    for cat, cnt in report["category_counts"].items():
        print(f"   • {cat}: {cnt} questions")
    print(f"Master JSON saved to: {MASTER_JSON}")
    print(f"Import report saved to: {REPORT_FILE}")
    print("==================================================")

if __name__ == '__main__':
    run_import()
